#!/usr/bin/env python3
"""Convertit « Coto JLPT Grammar and Vocab Mega List.xlsx » en fichiers de données pour l'appli.

Usage : python3 outils/import_coto.py chemin/vers/fichier.xlsx [N1 N2 ...]
Nécessite openpyxl (pip install openpyxl).
Produit web/data/<niveau>-vocab.js, data/<niveau>-kanji.js, data/<niveau>-grammaire-liste.js
"""
import json, re, sys, pathlib
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / "web" / "data"

def kana_type(s):
    s = s.replace(".", "").replace("-", "")
    if s and all("゠" <= c <= "ヿ" for c in s):
        return "on"
    return "kun"

def rows(ws, start):
    for r in ws.iter_rows(min_row=start, values_only=True):
        if not r or not (r[0] or r[1]):
            continue
        yield [("" if c is None else str(c).strip()) for c in r[:3]]

def header_row(ws):
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), start=1):
        if r and r[0] in ("Kanji", "Grammar Point"):
            return i
    return 2

ORDER = ["N5", "N4", "N3", "N2", "N1"]
# Corrections de sens anglais cassés dans la liste source (« #NAME? » d’Excel)
EN_FIX = {"歳": "age, -years-old", "ヶ月": "month(s) (counter)"}

def lower_words(wb, lvl, sheet):
    """Mots/kanji déjà présents dans les niveaux inférieurs (la liste Coto contient des doublons)."""
    out = set()
    for low in ORDER[:ORDER.index(lvl)]:
        ws = wb[f"{low} {sheet}"]
        for r in ws.iter_rows(values_only=True):
            if r and r[0] and str(r[0]).strip() not in ("Kanji",) and not str(r[0]).startswith("="):
                out.add(str(r[0]).strip())
    return out

def main(path, levels):
    wb = openpyxl.load_workbook(path, read_only=True)
    for lvl in levels:
        L = lvl.lower()
        lowv, lowk = lower_words(wb, lvl, "Vocab"), lower_words(wb, lvl, "Kanji")
        # Vocabulaire
        ws = wb[f"{lvl} Vocab"]
        seen, vocab = set(), []
        for w, r, m in rows(ws, header_row(ws) + 1):
            if w.startswith("=") or r.startswith("="):
                continue
            w = w or r
            key = f"{w}|{r}"
            if key in seen or w in lowv:
                continue
            seen.add(key)
            if m.startswith("#") or not m:
                m = EN_FIX.get(w, m)
            vocab.append({"w": w, "r": r, "m": m})
        (DATA / f"{L}-vocab.js").write_text(
            f"window.DATA=window.DATA||{{}};window.DATA['{lvl}-vocab']=" + json.dumps(vocab, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
        # Kanji
        ws = wb[f"{lvl} Kanji"]
        kanji, seen = [], set()
        for k, p, m in rows(ws, header_row(ws) + 1):
            if not k or k.startswith("=") or k in seen or k in lowk:
                continue
            seen.add(k)
            parts = [x.strip() for x in re.split(r"[・、,]", p) if x.strip()]
            kanji.append({"k": k, "on": [x for x in parts if kana_type(x) == "on"],
                          "kun": [x for x in parts if kana_type(x) == "kun"], "m": m})
        (DATA / f"{L}-kanji.js").write_text(
            f"window.DATA=window.DATA||{{}};window.DATA['{lvl}-kanji']=" + json.dumps(kanji, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
        # Liste des points de grammaire (noms seulement)
        ws = wb[f"{lvl} Grammar"]
        gl = []
        for r in ws.iter_rows(min_row=header_row(ws) + 1, values_only=True):
            if r and r[0] and not str(r[0]).startswith("="):
                if str(r[0]).strip() not in gl:
                    gl.append(str(r[0]).strip())
        (DATA / f"{L}-grammaire-liste.js").write_text(
            f"window.DATA=window.DATA||{{}};window.DATA['{lvl}-grammaire-liste']=" + json.dumps(gl, ensure_ascii=False) + ";\n", encoding="utf-8")
        print(f"{lvl}: {len(vocab)} mots, {len(kanji)} kanji, {len(gl)} points de grammaire")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2:] or ["N1"])
